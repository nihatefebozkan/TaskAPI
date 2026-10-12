from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.6)
sec.left_margin = sec.right_margin = Cm(1.9)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)

def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t, 2)
def bullets(items):
    for t in items: doc.add_paragraph(t, 'List Bullet')
def code(t):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    r = par.add_run(t)
    r.font.name, r.font.size = 'Consolas', Pt(9)
def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)
def questions(items):
    h('Kendini sınama')
    for i,t in enumerate(items, 1): p(f'{i}. {t}')
    p('Notlarım: __________________________________________________')
def table(headers, rows):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.style = 'Table Grid'
    for c,t in zip(tb.rows[0].cells, headers):
        c.text=t
        for r in c.paragraphs[0].runs: r.bold=True
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DCE6F1'); c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        for c,t in zip(tb.add_row().cells,row): c.text=t
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    tb._tbl.tblPr.append(borders)
    for row in tb.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(4)
                par.paragraph_format.space_before=Pt(4)
                for r in par.runs: r.font.size=Pt(10)
    repeat=OxmlElement('w:tblHeader'); tb.rows[0]._tr.get_or_add_trPr().append(repeat)
    p('')

doc.add_heading('TaskAPI Auth Çalışma Kitabı', 0)
p('Kimlik doğrulama ve yetkilendirme için sekiz fazlı öğrenme programı')
p('9 Ekim 2026')
h('Bu dokümanla ne yapacağız')
p('Hedefimiz, auth kavramlarını açıklayabilmek ve TaskAPI’de giriş yapan kullanıcının yalnızca kendi görevlerini yönetebildiği bir akış kurmak. Basic Auth, session, token ve JWT’yi karşılaştıracağız; ardından OAuth 2.0 ve OpenID Connect ile harici giriş mantığını öğreneceğiz.')
p('Bu bir çalışma kitabıdır. Her fazı okuyup örnek akışı kendi cümlelerinle anlat, alıştırmayı yap ve soruları cevapla. Cevap anahtarına en son bak. Uygulama görevleri gelecekte yapacağımız işlerdir; tamamlanmış proje özellikleri değildir.')
h('Başlangıç durumumuz')
p('TaskAPI’de kullanıcı kaydı, şifre hashleme, kullanıcı adı ve şifre kontrolü, görev CRUD işlemleri ve test dosyaları var. Login bool sonuç döndürüyor. JWT üretimi, JWT doğrulaması ve görevlerde kullanıcı erişim kontrolü henüz kurulmamış. Şifre hashleme altyapısı ile JWT imzalama farklı sorumluluklardır.')
table(['Faz','Konu','Somut çıktı'],[
('1','HTTP ve auth temelleri','Kimlik ve yetki ayrımı'),('2','Basic Authentication','Örnek Basic isteği'),('3','Session ve cookie','Oturum akışı çizimi'),('4','Token ve JWT','Claim ve doğrulama açıklaması'),('5','TaskAPI JWT girişi','Korumalı endpoint'),('6','Sahiplik ve roller','Kullanıcıya özel görevler'),('7','Yaşam döngüsü ve testler','Yenileme ve çıkış kuralları'),('8','OAuth 2.0 ve OIDC','Harici giriş akışı')])
p('Çalışma düzeni: Ders → örnek → kendi açıklaman → uygulama → doğrulama. Takvimi ezbere değil, her fazın tamamlanma ölçütüne göre ilerlet.')

page('Faz 1 HTTP ve auth temelleri')
p('Amaç: Bir login isteğinin neden sonraki isteklere otomatik kimlik sağlamadığını anlamak.')
h('İstek nasıl okunur')
p('HTTP isteğinde metot yapılmak istenen işlemi, URL hedef kaynağı, header ek bilgileri, body ise gönderilen veriyi taşır. GET /api/task/42 içindeki 42 görev ID’sidir; isteği gönderen kullanıcının kimliğini göstermez. JSON verinin biçimidir, kimlik doğrulama yöntemi değildir.')
code('GET /api/task/42\nAuthorization: Bearer <token>\nAccept: application/json')
p('Authorization header’ı kimlik bilgisini taşır. Accept istemcinin istediği cevap biçimini belirtir. Content-Type ise gönderilen gövdenin biçimini açıklar. Cookie de bir HTTP header’ı üzerinden taşınır.')
h('Login ile sonraki istek arasındaki boşluk')
p('HTTP istekleri kendi başlarına önceki kullanıcı girişini hatırlamaz. Sunucu oturum veya token gibi bir mekanizma kurarak bu ilişkiyi sağlar. Aynı ağ bağlantısının kullanılması kullanıcı kimliği yerine geçmez. TaskAPI’de login true döndürdüğünde şifre kontrolü tamamlanır; fakat sonraki görev isteğine taşınacak kimlik henüz oluşturulmaz.')
code('Kullanıcı bilgileri → Login → Kimliği temsil eden bilgi\nYeni istek + bu bilgi → Doğrulama → Yetki kontrolü → İşlem')
h('Kimlik ve yetki')
p('Authentication, isteği gönderenin kim olduğunu belirler. Authorization, bu kişinin istenen işlemi yapıp yapamayacağını belirler. Ahmet’in giriş yapması, Mehmet’in görevini silme yetkisi vermez. Kontroller sunucuda yapılır; arayüzde bir düğmeyi gizlemek yeterli değildir.')
table(['Durum','Beklenen cevap'],[('Kimlik bilgisi yok veya geçersiz','401'),('Kimlik doğrulandı ama yetki yetersiz','403'),('İstek başarıyla işlendi','İşleme uygun 2xx kodu')])
h('Alıştırma ve tamamlanma ölçütü')
p('Bir login isteği ve bir görev okuma isteği yaz. Her birinde URL, metot, header ve body’yi işaretle. Ardından authentication ve authorization kontrollerinin nerede yapılacağını göster. Login’in neden tek başına yeterli olmadığını açıklayabiliyorsan bu faz tamamdır.')
questions(['Ahmet login oldu; sonraki korumalı istekte kimlik bilgisi göndermedi. API onu tanıyabilir mi?', 'Geçerli kimlikle admin gerektiren işleme gelen normal kullanıcıya hangi kod verilir?'])

page('Faz 2 Basic Authentication')
p('Amaç: Kullanıcı adı ve şifrenin her istekte nasıl taşınabildiğini öğrenmek. Basic Auth’u TaskAPI’nin JSON login endpoint’iyle karıştırmamak.')
h('Akışın mantığı')
p('İstemci username:password metnini Base64 ile kodlar ve Authorization header’ında Basic şemasıyla gönderir. Sunucu bilgileri çözümler, kullanıcıyı doğrular ve işlem yetkisini kontrol eder. Basic yöntemi kendi başına access token üretmez. [1]')
code('Örnek metin: ahmet:ornek\nBase64: YWhtZXQ6b3JuZWs=\n\nGET /api/task\nAuthorization: Basic YWhtZXQ6b3JuZWs=')
p('Bu bilgiler yalnızca öğrenme örneğidir. Base64 geri çözülebilen bir kodlamadır; gizlilik sağlamaz. Basic kimlik bilgileri HTTPS olmadan ağ üzerinde korunmaz. [1]')
h('Üç kavramı ayıralım')
bullets(['Kodlama: Veriyi başka bir gösterime çevirir. Base64 buna örnektir.', 'Şifreleme: Anahtar kullanarak veriyi gizler; uygun anahtarla geri açılabilir.', 'Hashleme: Parola saklama gibi amaçlarda tek yönlü türetme kullanılır. Parola doğrulama için uygun algoritma ve salt gerekir.'])
p('TaskAPI’nin şifre hashleme işlemi veritabanında şifreyi saklama konusudur. Basic header ise kimlik bilgisinin istekte taşınması konusudur. Birini kullanmak diğerinin görevini karşılamaz.')
h('JSON login ile farkı')
p('POST /api/auth/login gövdesinde kullanıcı adı ve şifre göndermek otomatik olarak Basic Auth değildir. Yöntem, kimlik bilgisinin biçimine ve sunucunun onu nasıl işlemesine göre belirlenir. Login sonrasında JWT döndürürsek sonraki isteklerde Bearer şeması kullanırız.')
h('TaskAPI çalışması')
p('Basic ve Bearer için iki örnek header yaz. Kimlik bilgisinin her istekte nasıl doğrulandığını karşılaştır. Bu fazda ana uygulamaya ikinci bir auth sistemi eklemek zorunlu değildir; ayrı küçük bir örnek yeterlidir.')
h('Tamamlanma ölçütü')
p('Basic isteğini açıklayabiliyor, Base64’ün şifreleme olmadığını gösterebiliyor ve şifre hashleme ile istek doğrulamasını ayırabiliyorsan ilerle.')
questions(['Basic header içinde kodlanan metnin biçimi nedir?', 'JSON body ile şifre gönderen her endpoint Basic Auth sayılır mı?'])

page('Faz 3 Session ve cookie')
p('Amaç: Sunucunun giriş durumunu takip etmesini ve cookie’nin bu akıştaki görevini anlamak.')
h('Sunucuda tutulan oturum')
p('Klasik session yaklaşımında login başarılı olduğunda sunucu bir oturum kaydı oluşturur. İstemciye tahmin edilmesi zor bir session ID verir. Tarayıcı bu kimliği cookie’de saklar ve uygun isteklere ekler. Sunucu ID’yi oturum kaydıyla eşleştirerek kullanıcıyı bulur.')
code('Login → Sunucuda oturum kaydı → Cookie ile session ID\nGörev isteği + cookie → Oturumu bul → Kullanıcıyı tanı')
p('Cookie taşıma ve saklama mekanizmasıdır; session ise oturum yönetimi yaklaşımıdır. Her cookie bir oturum kimliği taşımaz. Bir cookie başka tür veri veya korunan bir kimlik bileti de taşıyabilir.')
h('ASP.NET Core farkı')
p('ASP.NET Core cookie authentication varsayılan olarak korunan bir kimlik bileti oluşturabilir; her istekte ayrı bir sunucu session tablosu araması zorunlu değildir. Ayrıca genel amaçlı session state, tek başına authentication sistemi değildir. Bu üç kavramı aynı şey sanmamalıyız. [2]')
h('Cookie özellikleri')
bullets(['HttpOnly: Cookie’nin JavaScript ile okunmasını engeller; tüm XSS etkilerini ortadan kaldırmaz.', 'Secure: Cookie’nin HTTPS üzerinden gönderilmesini sağlar.', 'SameSite: Cookie’nin farklı site bağlamlarında gönderilmesine ilişkin kuralları etkiler.'])
p('Tarayıcı cookie’yi otomatik gönderebildiği için cookie kullanan işlemlerde CSRF düşünülmelidir. Domain, path ve tarayıcı politikaları da hangi isteğe cookie eklendiğini etkiler. Cookie tabanlı auth, otomatik olarak tüm güvenlik sorunlarını çözmez.')
h('Çıkış ve ölçekleme')
p('Sunucu oturumunu iptal etmek ile tarayıcıdaki cookie’yi silmek farklı işlemlerdir. Klasik merkezi session kaydı iptal edildiğinde sonraki istek reddedilebilir. Birden fazla sunucu kullanıldığında oturumların nasıl paylaşılacağı ayrıca tasarlanır.')
h('Alıştırma ve tamamlanma ölçütü')
p('Tarayıcı ve sunucu için iki sütun çiz. Session ID, kullanıcı kaydı, cookie ve oturum bitiş zamanını uygun yere yerleştir. Sonra logout sırasında hangi parçaların değişeceğini anlat. Taşıma mekanizmasıyla oturum yönetimini ayırabiliyorsan bu faz tamamdır.')
questions(['Cookie ile session aynı şey midir?', 'HttpOnly işaretli cookie, JavaScript üzerinden doğrudan okunabilir mi?'])

page('Faz 4 Token ve JWT')
p('Amaç: Token yaklaşımını öğrenmek ve JWT’nin bir token biçimi olduğunu anlamak.')
h('Bearer ve token türleri')
p('Bearer kullanımında token’ı elinde bulunduran kişi genellikle onu sunarak erişebilir. Bu yüzden token gizli tutulmalıdır. Opaque token, istemcinin içeriğinden anlam çıkaramadığı bir değerdir; sunucu onu bir kayıt veya doğrulama hizmetiyle çözebilir. JWT ise claim’leri yapılandırılmış biçimde taşır.')
h('JWT parçaları')
code('header.payload.signature\n\nÖrnek claim’ler\n{"sub":"42", "iss":"taskapi-auth", "aud":"taskapi",\n "role":"user", "exp":1800000000}')
p('Burada yaygın, imzalı kompakt JWT biçimini öğreniyoruz. Header algoritma bilgisini, payload claim’leri, signature bütünlük kontrolünü taşır. Header ve payload Base64url ile kodlanır; imza içeriği gizlemez. JWT’nin şifrelenmiş biçimleri de vardır.')
table(['Claim','Anlamı'],[('sub','Token’ın temsil ettiği özne'),('iss','Token’ı üreten taraf'),('aud','Hedef alıcı veya API'),('exp','Sona erme zamanı'),('iat','Üretilme zamanı'),('nbf','Bu zamandan önce kullanılamaz')])
p('Zaman claim’leri Unix zamanına göre saniye değerleri taşır. Örnek exp değeri yalnızca biçimi gösterir. Claim’lerin zorunluluğunu uygulama ve ilgili protokol belirler; tüm JWT’lerde bu alanların tamamı bulunmak zorunda değildir.')
h('Üç ayrı işlem')
p('Decode içeriği okur. Validate imza ve geçerlilik koşullarını değerlendirir. Authorize ise doğrulanmış kimliğin işlem yetkisini kontrol eder. Algoritma ve anahtar kabulü sunucuda sınırlandırılır; token’ın kendi beyanı tek başına yeterli olmaz. [3]')
p('HMAC yaklaşımında imzalayan ve doğrulayan taraf ortak sır kullanır. Asimetrik yaklaşımda özel anahtar imzalar, açık anahtar doğrular. Kriptografik işlemleri kendimiz yazmak yerine uygun kütüphanelere bırakırız.')
h('Alıştırma ve tamamlanma ölçütü')
p('Claim tablosunu kapatıp her alanı açıklamayı dene. Bir rol değişikliğinin eski token’ın içeriğini neden değiştirmediğini anlat. Okuma ile doğrulama farkını açıklayabiliyorsan ilerle.')
questions(['JWT içindeki role değiştirildiğinde eski imza neden geçersiz olur?', 'JWT’ye kullanıcı şifresi koymak uygun mudur?'])

page('Faz 5 TaskAPI içinde JWT girişi')
p('Amaç: Login’den access token döndürmek ve görev endpoint’lerinde doğrulanmış kimlik oluşturmak. Bu bölüm uygulama planıdır; henüz projeye eklenmiş kod değildir.')
h('Sorumlulukları belirle')
bullets(['AccountService: Kullanıcı bilgilerini doğrular.', 'Token servisi: Gerekli claim’lerle token üretir.', 'JWT authentication katmanı: Gelen token’ı doğrular.', 'Controller ve policy: Endpoint erişimini belirler.'])
h('Uygulama sırası')
bullets(['Issuer, audience, access token süresi ve imzalama anahtarını yapılandır.', 'Anahtarı kaynak koda koyma; geliştirme ve üretim için uygun sır yönetimi kullan.', 'Token üretme servisi ve login cevap DTO’su oluştur.', 'Login başarı durumunda access token ve bitiş bilgisini döndür.', 'AddAuthentication ve AddJwtBearer ile doğrulama yapılandırmasını ekle.', 'UseAuthentication çağrısını UseAuthorization öncesinde çalıştır.', 'TaskController üzerinde Authorize kullan; register ve login erişimini açık tut.', 'UI’den sonraki isteklerde Bearer token gönder.'])
p('JWT bearer doğrulamasında imza, issuer, audience ve süre kontrolleri yapılandırılmalıdır. Kimlik doğrulama handler’ı doğrulanmış kullanıcı bilgisini oluşturur. Yalnızca UseAuthorization eklemek JWT doğrulamasını kurmaz. [4]')
code('POST /api/auth/login\n→ {"accessToken":"<token>", "expiresAt":"<UTC zamanı>"}\n\nGET /api/task\nAuthorization: Bearer <token>')
h('Doğrulama çalışması')
table(['Deneme','Beklenti'],[('Doğru giriş bilgileri','Token cevabı'),('Yanlış giriş bilgileri','401 ve tutarlı hata'),('Korumalı istekte token yok','401'),('Değiştirilmiş veya süresi dolmuş token','401'),('Geçerli ve uygun token','Endpoint çalışır')])
p('Mevcut login servisindeki genel exception yaklaşımını da ele alacağız: Yanlış şifre, beklenen kimlik doğrulama hatası olarak yönetilmeli; beklenmeyen sistem hatasıyla aynı davranışa düşmemelidir.')
h('Tamamlanma ölçütü')
p('İstemci giriş yapıp token alır; korumalı endpoint yalnızca doğrulanan token ile çalışır. Bu aşama tek başına kullanıcıların görevlerini birbirinden ayırmaz; sahiplik Faz 6’da eklenir.')
questions(['Token üretmek ile gelen token’ı doğrulamak aynı sorumluluk mudur?', 'Authorize eklemek tek başına kullanıcıya özel görev filtrelemesi yapar mı?'])

page('Faz 6 Kullanıcı sahipliği ve roller')
p('Amaç: Ahmet’in Mehmet’in görevlerini okuyamamasını veya değiştirememesini sağlamak. Kimliği doğrulamak bu kuralı kendiliğinden oluşturmaz.')
h('Sahipliği veriye bağla')
p('Task modeline kullanıcı sahipliğini gösteren bir alan eklenir. Görev oluştururken sahip, istemcinin gönderdiği kullanıcı ID’sinden değil doğrulanmış kullanıcı kimliğinden belirlenir. Böylece istemci başka bir kullanıcı adına görev oluşturamaz.')
code('Doğrulanmış sub → Kullanıcı ID’si\nYeni görev → OwnerUserId = doğrulanmış kullanıcı ID’si\nListeleme → Yalnızca bu kullanıcının görevleri')
h('Her işlemde kontrol')
bullets(['Listeleme: Kullanıcının görevleri filtrelenir.', 'Tek görev okuma: ID ile birlikte sahiplik kontrol edilir.', 'Güncelleme ve silme: Aynı sahiplik kuralı uygulanır.', 'İstemcinin gönderdiği ID: Yetki kanıtı olarak kabul edilmez.'])
p('Başka kullanıcıya ait görev istenirse 403 veya kaynağın varlığını gizlemek için 404 tercih edilebilir. Bu projede davranışı belirleyip tutarlı uygulayacağız. Kullanıcının görev listesinin boş olması ise normal bir sonuçtur.')
h('Rol ve policy')
p('Rol, kullanıcıyı user veya admin gibi bir gruba bağlar. Policy bir erişim kuralını tanımlar. Admin rolünün hangi işlemleri yapabileceği açıkça belirlenmelidir; rolün token’da yazması kendi başına iş kuralı oluşturmaz. Görev sahibi olma kontrolü de yalnızca rol kontrolüyle karşılanmaz.')
p('UI’de sil düğmesini gizlemek kullanıcı deneyimine yardım eder; güvenlik kararı API’de verilir. İstemci elle HTTP isteği gönderebilir. JWT’nin geçerli olması, istenen görev üzerinde yetki olduğunu göstermez.')
h('İki kullanıcıyla alıştırma')
p('Ahmet ve Mehmet hesabı oluştur. Her biri bir görev oluştursun. Ahmet’in token’ıyla Mehmet’in görevini okumayı, güncellemeyi ve silmeyi dene. Sonra kendi görevi üzerinde aynı işlemleri dene. Her endpoint’in aynı sahiplik kuralını uyguladığını doğrula.')
h('Tamamlanma ölçütü')
p('Kullanıcılar birbirlerinin görevlerine erişemez. Sahiplik sunucuda atanır. Varsa admin istisnaları açıkça tanımlanır ve test edilir.')
questions(['Görev oluştururken body’de gönderilen userId neden doğrudan kabul edilmemeli?', 'Geçerli token, tüm görevleri silme yetkisi verir mi?'])

page('Faz 7 Token yaşam döngüsü ve güvenlik')
p('Amaç: Token’ın üretilmesinden sonra yenileme, çıkış ve değişen kullanıcı durumu için kurallar oluşturmak.')
h('Access ve refresh token')
p('Access token API erişiminde kullanılır. Refresh token yeni access token istemek için kullanılır; görev endpoint’ine gönderilmez. Süreler uygulamanın riskine ve kullanımına göre seçilir. Kısa access token süresi tek başına tüm iptal sorunlarını çözmez.')
p('TaskAPI için tasarım çalışması: Refresh token’ı rastgele üret, sunucuda hash’iyle ve kullanıcı bağlantısıyla sakla; bitiş, iptal ve token ailesi bilgilerini yönet. Yenileme sırasında eski token’ı tüketip yenisini üret. Aynı eski token’ın tekrar kullanımında ilgili oturumun nasıl reddedileceğini belirle.')
p('Rotasyonun eşzamanlı isteklerde de doğru işlemesi gerekir. OAuth güvenlik rehberi, public client refresh token’larında rotasyon veya gönderene bağlama gibi tekrar kullanım tespit önlemlerini ele alır. [5]')
h('Çıkışın gerçek anlamı')
p('İstemciden token’ı silmek sunucudaki kopyaları kendiliğinden iptal etmez. Refresh token’ın iptali yenilemeyi durdurur; daha önce verilmiş access token ayrı bir iptal kontrolü yoksa süresi dolana kadar kabul edilebilir. Anında iptal gerekiyorsa sunucu tarafı kontrol ve bunun maliyeti tasarlanır.')
h('Saklama ve saldırı yüzeyi')
bullets(['JavaScript’in okuyabildiği depolardaki token’lar XSS ile ele geçirilebilir.', 'HttpOnly cookie doğrudan okumayı sınırlar; cookie akışında CSRF önlemleri ayrıca gerekir.', 'Token ve şifreleri loglara yazma; ağ iletişiminde HTTPS kullan.', 'Login denemelerini sınırlandır; kullanıcı varlığını gereksiz yere açığa çıkarma.', 'Parola veya rol değişiminde eski oturumların davranışını belirle.'])
h('Test çalışması')
p('Birim testleriyle iş kurallarını; HTTP entegrasyon testleriyle auth middleware’i, 401/403 cevaplarını ve endpoint erişimini doğrula. Yanlış audience, yanlış issuer, bozuk imza, süresi dolmuş token ve iptal edilmiş refresh token senaryolarını ekle. Çıkış ve rol değişimi için seçilen davranışı da test et.')
h('Tamamlanma ölçütü')
p('Giriş, yenileme ve çıkış akışları çalışır. Tek cihaz ve tüm cihazlardan çıkış davranışı tanımlıdır. Rol veya parola değişince mevcut erişimin ne zaman sonlanacağı açıklanabilir ve testlerle doğrulanır.')
questions(['Refresh token iptal edilince access token her zaman anında geçersiz olur mu?', 'Rotasyondan sonra eski refresh token tekrar gönderilirse ne incelenmelidir?'])

page('Faz 8 OAuth 2 ve OpenID Connect')
p('Amaç: Bir uygulamaya erişim izni vermek ile bir kullanıcıyı harici sağlayıcı üzerinden tanımak arasındaki farkı öğrenmek.')
h('OAuth 2 rollerini öğren')
p('Resource owner kaynağın sahibi, client erişim isteyen uygulama, authorization server izin ve token işlemlerini yöneten taraf, resource server ise korunan API’dir. Scope istenen erişim kapsamını belirtir. OAuth tek başına kullanıcı girişini tanımlayan bir kimlik protokolü değildir.')
p('Örnek: Bir uygulamaya Drive dosyalarını okuma izni vermek erişim yetkilendirmesidir. Bu izin uygulamaya kullanıcının Google şifresini vermek anlamına gelmez.')
h('Authorization Code ve PKCE')
code('Uygulama → Sağlayıcının giriş ve izin ekranı\nSağlayıcı → Uygulamaya authorization code\nUygulama → Code ve PKCE doğrulamasıyla token isteği\nAccess token → İzin verilen API’ye erişim')
p('PKCE, authorization code’u akışı başlatan istemciye bağlamaya yardımcı olur. Redirect URI, akış durumu ve sağlayıcı doğrulamaları da önemlidir. Tarayıcı uygulamasına kalıcı client secret koymak onu gizli tutmaz. Kullanıcısız hizmet erişimi için Client Credentials ayrı bir akıştır. [5]')
h('OIDC ne ekler')
p('OpenID Connect, OAuth 2.0 üzerine kullanıcı kimliğini doğrulama katmanı ekler. openid scope’u bu talebi belirtir. ID token giriş yapan kullanıcı hakkında client’a bilgi verir. Access token korunan API’ye erişmek içindir; ID token’ı API access token’ı yerine kullanmayız. [6]')
p('state akışı başlatan istekle geri dönüşü ilişkilendirmede kullanılır. nonce gönderilmişse ID token’da doğrulanır ve tekrar kullanım risklerine karşı akış bağlantısını güçlendirir. Kullanılan akış ve kütüphanenin doğrulama kurallarını takip ederiz.')
h('TaskAPI çalışması')
p('Önce yerel JWT akışını tamamla. Sonra bir harici kimlik sağlayıcı seçip OIDC girişini ayrı bir uygulama adımı olarak incele. Kullanıcıyı yalnızca görünen adla eşleştirme; sağlayıcının issuer ve subject kimliğini birlikte değerlendir. Sağlayıcının token’ının hangi API için geçerli olduğunu kontrol et.')
h('Tamamlanma ölçütü')
p('OAuth ve OIDC’nin görevlerini ayırabilir, authorization code akışını çizebilir ve ID token ile access token’ın hedefini açıklayabilirsin. Harici sağlayıcı entegrasyonu ikinci uygulama hedefimizdir.')
questions(['Google ile girişte kullanıcı kimliğini standart biçimde sağlayan katman hangisidir?', 'ID token neden doğrudan TaskAPI access token’ı yerine kullanılmamalıdır?'])

page('Cevap anahtarı ve tekrar')
p('Önce kendi cevaplarını yaz. Aşağıdaki açıklamalarla karşılaştır; ezber yerine nedenini anlatmayı hedefle.')
answers=[
('Faz 1', '1. İstekte kimlik bilgisi yoksa önceki login’e bakarak Ahmet olduğunu varsayamayız; korumalı endpoint 401 döndürür. 2. Kimlik doğrulanmış ama gerekli yetki yoksa 403.'),
('Faz 2', '1. username:password metni Base64 ile kodlanır. 2. Hayır; JSON login ayrı bir akıştır, Basic şemasıyla aynı değildir.'),
('Faz 3', '1. Hayır; cookie veriyi taşır, session oturumu yönetir. 2. HttpOnly cookie JavaScript ile doğrudan okunamaz; bu tüm XSS etkilerini engellemez.'),
('Faz 4', '1. İmza imzalanan içeriğe bağlıdır; değişen payload eski imzayla doğrulanmaz. 2. Hayır; yaygın imzalı JWT’nin payload’ı okunabilir.'),
('Faz 5', '1. Hayır; üretim login tarafında, doğrulama gelen API isteğinde yapılır. 2. Hayır; Authorize kimlik veya tanımlı policy kontrolüdür, görev sahipliği ayrıca uygulanır.'),
('Faz 6', '1. İstemci başka bir kullanıcı ID’si gönderebilir; sahip doğrulanmış kimlikten atanmalıdır. 2. Hayır; kaynak ve işlem yetkisi ayrıca kontrol edilir.'),
('Faz 7', '1. Hayır; access token için ayrıca iptal kontrolü yoksa süre sonuna kadar erişim sürebilir. 2. Olası token çalınması, tekrar kullanım ve eşzamanlı istek durumu; seçilen oturum iptal politikası uygulanır.'),
('Faz 8', '1. OpenID Connect. 2. ID token client’a kimlik bilgisi verir; API erişim token’ıyla hedefi ve doğrulama kuralları farklıdır.')]
for title,text in answers: h(title); p(text)
h('Kendi cümlelerinle tekrar et')
p('Bir dakika içinde şu akışı anlat: Kullanıcı login olur, kimliği doğrulanır, access token alır, görev isteğine ekler, API token’ı doğrular, görev sahipliğini kontrol eder ve cevap verir. Ardından token süresi dolunca ve kullanıcı çıkış yapınca ne olacağını açıkla.')

page('Uygulama planı ve kaynaklar')
h('İlk uygulama hedefi')
p('Faz 1–4 kavramları hazırlar. Faz 5–6 sonunda kullanıcı giriş yapar ve yalnızca kendi görevlerini yönetir. Faz 7 yaşam döngüsünü tamamlar. Faz 8 harici giriş seçeneklerini açar. Her yöntemi aynı projede kullanmak zorunda değiliz.')
bullets(['Login cevap modeli ve token servisi hazır.', 'JWT bearer doğrulaması ve korumalı endpoint hazır.', 'Görev sahipliği sunucuda atanıyor ve kontrol ediliyor.', 'Yanlış kimlik ile yetersiz yetki ayrılıyor.', 'Token yenileme ve çıkış davranışı belirlenmiş.', 'Auth ve sahiplik sınırları testlerle doğrulanmış.'])
h('Sık karıştırılan noktalar')
p('JWT bir token biçimidir. Bearer bir sunma yaklaşımıdır. Cookie bir taşıma mekanizmasıdır. Session oturum yönetimidir. OAuth erişim izni çerçevesidir. OIDC kullanıcı kimliğini doğrulama katmanıdır. Bu kavramlar bazı sistemlerde birlikte kullanılır.')
h('Resmi kaynaklar')
p('Kaynaklar 9 Ekim 2026 tarihinde kontrol edildi. İlgili fazdaki köşeli parantezli numara aşağıdaki kaynağa karşılık gelir. Ayrıntılı uygulamada kullanılan framework sürümünün belgelerini takip et.')
sources=[
('[1] Basic Authentication standardı', 'https://www.rfc-editor.org/rfc/rfc7617.html'),
('[2] ASP.NET Core cookie authentication', 'https://learn.microsoft.com/en-us/aspnet/core/security/authentication/cookie?view=aspnetcore-10.0'),
('[3] JWT güvenlik uygulamaları', 'https://www.rfc-editor.org/rfc/rfc8725.html'),
('[4] ASP.NET Core JWT bearer yapılandırması', 'https://learn.microsoft.com/en-us/aspnet/core/security/authentication/configure-jwt-bearer-authentication?view=aspnetcore-10.0'),
('[5] OAuth 2.0 güvenlik uygulamaları', 'https://www.rfc-editor.org/rfc/rfc9700.html'),
('[6] OpenID Connect Core', 'https://openid.net/specs/openid-connect-core-1_0.html')]
for label,url in sources:
    p(label)
    par=doc.add_paragraph(); r=par.add_run(url); r.font.size=Pt(9)
h('Bir sonraki çalışma notum')
p('En çok zorlandığım kavram: ___________________________________')
p('Tekrar edeceğim faz: _________________________________________')
p('TaskAPI’de deneyeceğim ilk işlem: ______________________________')

footer=sec.footer.paragraphs[0]
footer.alignment=2
r=footer.add_run('TaskAPI Auth Çalışma Kitabı  |  '); r.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='TaskAPI Auth Çalışma Kitabı'
doc.core_properties.subject='Sekiz fazlı kimlik doğrulama ve yetkilendirme çalışma programı'
doc.core_properties.author=''
out=ROOT/'TaskAPI_Auth_Calisma_Kitabi.docx'
doc.save(out)
print(out)
